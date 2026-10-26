from django.shortcuts import render, get_object_or_404

# Create your views here.
from programs.models import Program
from .models import TestPoint

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def test_matrix(request, program_id):
    program = get_object_or_404(
        Program,
        pk=program_id,
    )

    test_points = (
        TestPoint.objects
        .filter(
            test__engineering_requirement__program=program
        )
        .select_related(
            'test',
            'test__engineering_requirement',
            'test__engineering_requirement__program',
        )
        .order_by('test__test_number', 'test_point_number')
    )

    context = {
        'program': program,
        'test_points': test_points,
    }

    return render(
        request,
        'testing/test_matrix.html',
        context,
    )

def export_test_matrix(request, program_id):

    # 1. Get selected aircraft program
    program = get_object_or_404(
        Program,
        pk=program_id
    )

    # 2. Retrieve test points for this program
    test_points = (
        TestPoint.objects
        .filter(
            test__engineering_requirement__program=program
        )
        .select_related(
            'test',
            'test__engineering_requirement'
        )
        .prefetch_related('conditions')
        .order_by(
            'test__test_number',
            'test_point_number'
        )
    )

    # 3. Create Excel workbook
    workbook = Workbook()
    worksheet = workbook.create_sheet('Flight Test Matrix')
    worksheet.remove(workbook.worksheets[0])
    worksheet.title = "Flight Test Matrix"

    # 4. Define columns
    headers = [
        "Test Number",
        "Test Title",
        "Requirement",
        "Test Point",
        "Weight",
        "CG",
        "Target Speed",
        "Target Altitude",
        "Repetitions",
        "Conditions",
        "Pass/Fail Criteria",
        "Status",
        "Remarks",
    ]

    worksheet.append(headers)

    # 5. Populate data
    for point in test_points:

        conditions = "; ".join(
            f"{condition.name}: {condition.value} "
            f"{condition.unit}".strip()
            for condition in point.conditions.all()
        )

        worksheet.append([
            point.test.test_number,
            point.test.test_title,
            point.test.engineering_requirement.document_number,
            point.test_point_number,
            point.weight_category,
            point.cg_category,
            point.target_speed,
            point.target_altitude,
            point.repetition_required,
            conditions,
            point.pass_fail_criteria,
            point.status,
            point.remarks,
        ])

    # 6. Format header
    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            start_color="17365D",
            end_color="17365D",
            fill_type="solid"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # 7. Set column widths
    for column in worksheet.columns:
        column_letter = get_column_letter(
            column[0].column
        )
        worksheet.column_dimensions[column_letter].width = 22

    # 8. Improve usability
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # 9. Prepare HTTP response
    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        f'attachment; filename="FTM_{program.pk}.xlsx"'
    )

    # 10. Write workbook to response
    workbook.save(response)

    return response