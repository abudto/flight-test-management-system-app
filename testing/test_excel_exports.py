"""Read actual workbook responses, preserving current export behavior.

Known defect: openpyxl interprets user text beginning '=' as a formula. The
expected-failure test describes the required literal-text behavior; remove its
marker when the exporter is fixed. It must then pass without data_only=True,
which would hide formulas behind cached results.
"""

from io import BytesIO
from unittest import expectedFailure

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from programs.models import Program
from testing.test_support import MatrixDataMixin


class ExcelMatrixExportTests(MatrixDataMixin, TestCase):
    HEADERS = [
        "Test Number", "Test Title", "Requirement", "Test Point", "Weight", "CG",
        "Target Speed", "Target Altitude", "Repetitions", "Conditions",
        "Pass/Fail Criteria", "Status", "Remarks",
    ]

    def export(self, program=None):
        return self.client.get(reverse(
            "export_test_matrix", args=[(program or self.program_a).pk],
        ))

    def worksheet(self, program=None):
        response = self.export(program)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Flight Test Matrix"])
        return workbook.active

    def test_response_type_and_attachment_filename(self):
        response = self.export()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(response["Content-Disposition"],
                         f'attachment; filename="FTM_{self.program_a.pk}.xlsx"')

    def test_worksheet_headers_and_structure(self):
        sheet = self.worksheet()
        self.assertEqual(sheet.max_column, 13)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual([cell.value for cell in sheet[1]], self.HEADERS)

    def test_export_filters_requested_program(self):
        for program, number, title in [
            (self.program_a, "PT-01", "Alpha performance test"),
            (self.program_b, "BRAVO-ONLY", "Bravo private test"),
        ]:
            with self.subTest(program=program.name):
                sheet = self.worksheet(program)
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(sheet["D2"].value, number)
                self.assertEqual(sheet["B2"].value, title)

    def test_rows_follow_test_then_point_identifier_order(self):
        expected = self.add_ordering_points()
        sheet = self.worksheet()
        self.assertEqual(
            [(row[0], row[3]) for row in sheet.iter_rows(min_row=2, values_only=True)],
            [(point.test.test_number, point.test_point_number) for point in expected],
        )

    def test_targets_and_engineering_metadata(self):
        sheet = self.worksheet()
        self.assertEqual([cell.value for cell in sheet[2]][:9], [
            "TEST-01", "Alpha performance test", "SPEC-01", "PT-01", "heavy",
            "Forward", 147.5, 6500, 3,
        ])
        for address in ["G2", "H2", "I2"]:
            self.assertEqual(sheet[address].data_type, "n")
        self.assertEqual([cell.value for cell in sheet[2]][10:], [
            "Speed remains within tolerance", "PLANNED", "Alpha engineering remarks",
        ])

    def test_conditions_include_values_and_units_without_program_leakage(self):
        conditions = self.worksheet()["J2"].value
        self.assertEqual(set(conditions.split("; ")),
                         {"Temperature: 15 degC", "Configuration: Clean"})
        self.assertNotIn("Bravo-only", conditions)

    def test_point_without_conditions_has_blank_cell(self):
        self.point_a.conditions.all().delete()
        self.assertIsNone(self.worksheet()["J2"].value)

    def test_zero_targets_remain_numeric(self):
        self.point_a.target_speed = 0
        self.point_a.target_altitude = 0
        self.point_a.save()
        sheet = self.worksheet()
        for address in ["G2", "H2"]:
            self.assertEqual(sheet[address].value, 0)
            self.assertEqual(sheet[address].data_type, "n")

    def test_missing_targets_remain_blank(self):
        self.point_a.target_speed = None
        self.point_a.target_altitude = None
        self.point_a.save()
        sheet = self.worksheet()
        self.assertIsNone(sheet["G2"].value)
        self.assertIsNone(sheet["H2"].value)

    def test_empty_program_exports_headers_only(self):
        empty = Program.objects.create(name="Empty synthetic program")
        sheet = self.worksheet(empty)
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual([cell.value for cell in sheet[1]], self.HEADERS)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:M1")

    def test_missing_program_returns_404(self):
        missing_id = max(self.program_a.pk, self.program_b.pk) + 1
        self.assertEqual(self.client.get(reverse(
            "export_test_matrix", args=[missing_id],
        )).status_code, 404)

    def test_header_is_frozen_and_filter_covers_all_rows(self):
        self.add_ordering_points()
        sheet = self.worksheet()
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:M5")

    @expectedFailure
    def test_formula_like_remarks_are_preserved_as_literal_text(self):
        """Known defect: '=1+1' currently becomes an Excel formula cell."""
        self.point_a.remarks = "=1+1"
        self.point_a.save()
        cell = self.worksheet()["M2"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s", "Engineering remarks must remain literal text")
